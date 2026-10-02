import csv
from pathlib import Path

import openpyxl


def convert_xlsx_to_csv(
    xlsx_path: Path,
    csv_path: Path,
    sheet_name: str | None = None,
    delimiter: str = ";",
) -> None:
    """Converts a worksheet of an .xlsx workbook to a CSV file."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Worksheet '{sheet_name}' not found in {xlsx_path}. Available: {wb.sheetnames}")
        sheet = wb[sheet_name]
    else:
        sheet = wb.active

    if sheet is None:
        raise ValueError(f"No active worksheet found in {xlsx_path}")

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=delimiter)
        for row in sheet.iter_rows(values_only=True):
            # Skip completely empty rows
            if not any(cell is not None for cell in row):
                continue
            writer.writerow([cell if cell is not None else "" for cell in row])


def normalize_and_copy_csv(src_csv: Path, dst_csv: Path) -> None:
    """Copies a CSV file while ensuring UTF-8 encoding without BOM and LF line endings."""
    dst_csv.parent.mkdir(parents=True, exist_ok=True)
    content = src_csv.read_bytes()
    # Remove UTF-8 BOM if present
    if content.startswith(b"\xef\xbb\xbf"):
        content = content[3:]
    # Convert CRLF to LF
    text = content.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
    dst_csv.write_text(text, encoding="utf-8")


def prepare_datasets(src_dir: Path, dst_dir: Path) -> list[Path]:
    """
    Scans the use cases source directory (src_dir) and prepares all target files
    in the destination directory (dst_dir).
    """
    src_dir = Path(src_dir).resolve()
    dst_dir = Path(dst_dir).resolve()
    dst_dir.mkdir(parents=True, exist_ok=True)

    mappings = [
        # GDP Use Case
        (
            src_dir / "gdp" / "Mapping Table.xlsx",
            dst_dir / "continent_mapping.csv",
            "xlsx",
            None,
        ),
        (
            src_dir / "gdp" / "life-expectancy-vs-gdp-per-capita - cleaned.csv",
            dst_dir / "life_expectancy.csv",
            "csv",
            None,
        ),
        # Superstore Use Case
        (
            src_dir / "superstore" / "Nomenclature.xlsx",
            dst_dir / "nomenclature.csv",
            "xlsx",
            None,
        ),
        (
            src_dir / "superstore" / "Sample - EU Superstore_Migrated Data - Part 1.csv",
            dst_dir / "superstore_part1.csv",
            "csv",
            None,
        ),
        (
            src_dir / "superstore" / "Sample - EU Superstore_Migrated Data - Part 2.csv",
            dst_dir / "superstore_part2.csv",
            "csv",
            None,
        ),
        # Allsales Use Case
        (
            src_dir / "allsales" / "Sales Workshop Files - version CTD - v2.xlsx",
            dst_dir / "allsales_part1.csv",
            "xlsx",
            "Sales table Part1 - 2025",
        ),
        (
            src_dir / "allsales" / "Sales Workshop Files - version CTD - v2.xlsx",
            dst_dir / "allsales_part2.csv",
            "xlsx",
            "Sales table Part2 - 2026",
        ),
        (
            src_dir / "allsales" / "Sales Workshop Files - version CTD - v2.xlsx",
            dst_dir / "allsales_team.csv",
            "xlsx",
            "Sales Team",
        ),
        (
            src_dir / "allsales" / "Sales Workshop Files - version CTD - v2.xlsx",
            dst_dir / "allsales_store.csv",
            "xlsx",
            "Store Locations",
        ),
    ]

    generated_files: list[Path] = []
    for src_path, dst_path, file_type, sheet_name in mappings:
        if not src_path.exists():
            raise FileNotFoundError(f"Source file not found: {src_path}")

        if file_type == "xlsx":
            sheet_info = f" (sheet: '{sheet_name}')" if sheet_name else ""
            print(f"-> Converting XLSX to CSV{sheet_info}: {src_path.name} => {dst_path.name}")
            convert_xlsx_to_csv(src_path, dst_path, sheet_name=sheet_name, delimiter=";")
        elif file_type == "csv":
            print(f"-> Normalizing CSV       : {src_path.name} => {dst_path.name}")
            normalize_and_copy_csv(src_path, dst_path)

        generated_files.append(dst_path)

    return generated_files
